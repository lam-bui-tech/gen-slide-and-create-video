"""
scripts/document_ingestion.py
Multi-format Document Ingestion Service for AI Lecture Studio
Extracts structured text, bullets, tables, embedded images, and speaker notes
from PPTX, PDF, DOCX, and Markdown files into DocumentContext.json conforming to schemas/.
"""

import os
import re
import json
import zipfile
import shutil
import xml.etree.ElementTree as ET
from typing import List, Dict, Any, Optional

from schemas.models import (
    DocumentContext,
    DocumentSection,
    DocumentTable,
    ExtractedImage,
    FileType,
)


class IngestionError(Exception):
    """Base exception for ingestion errors."""
    pass


class UnsupportedFormatError(IngestionError):
    """Raised when file extension is not supported."""
    pass


class FileTooLargeError(IngestionError):
    """Raised when file size exceeds maximum threshold."""
    pass


class ParseFailureError(IngestionError):
    """Raised when document parsing fails or file is corrupted."""
    pass


class DocumentIngestionService:
    def __init__(self, storage_base: str = "storage/projects", max_file_size_mb: int = 50):
        self.storage_base = storage_base
        self.max_file_size_bytes = max_file_size_mb * 1024 * 1024

    def ingest_document(self, file_path: str, project_id: str) -> DocumentContext:
        """
        Main entry point: validates file and parses into DocumentContext.
        Saves extracted images and document_context.json into storage/projects/{project_id}/.
        """
        if not os.path.exists(file_path):
            raise ParseFailureError(f"File not found: {file_path}")

        file_size = os.path.getsize(file_path)
        if file_size > self.max_file_size_bytes:
            raise FileTooLargeError(
                f"File size ({file_size / (1024*1024):.2f}MB) exceeds maximum limit of {self.max_file_size_bytes / (1024*1024):.0f}MB"
            )

        ext = os.path.splitext(file_path)[1].lower().lstrip(".")
        supported_exts = {
            "pptx": FileType.PPTX,
            "pdf": FileType.PDF,
            "docx": FileType.DOCX,
            "md": FileType.MD,
        }

        if ext not in supported_exts:
            raise UnsupportedFormatError(
                f"Unsupported file format: .{ext}. Supported formats are: pptx, pdf, docx, md."
            )

        file_type = supported_exts[ext]
        project_dir = os.path.join(self.storage_base, project_id)
        images_dir = os.path.join(project_dir, "images")
        os.makedirs(images_dir, exist_ok=True)

        try:
            if file_type == FileType.PPTX:
                doc_context = self._parse_pptx(file_path, project_id, images_dir)
            elif file_type == FileType.MD:
                doc_context = self._parse_markdown(file_path, project_id, images_dir)
            elif file_type == FileType.DOCX:
                doc_context = self._parse_docx(file_path, project_id, images_dir)
            elif file_type == FileType.PDF:
                doc_context = self._parse_pdf(file_path, project_id, images_dir)
            else:
                raise UnsupportedFormatError(f"No parser implemented for {ext}")
        except IngestionError:
            raise
        except Exception as e:
            raise ParseFailureError(f"Failed to parse {ext.upper()} document '{file_path}': {str(e)}") from e

        # Persist DocumentContext.json in project storage
        output_json_path = os.path.join(project_dir, "document_context.json")
        with open(output_json_path, "w", encoding="utf-8") as f:
            json.dump(doc_context.model_dump(), f, ensure_ascii=False, indent=2)

        return doc_context

    # ==========================================
    # 1. PPTX Parser
    # ==========================================
    def _parse_pptx(self, file_path: str, project_id: str, images_dir: str) -> DocumentContext:
        extracted_images: List[ExtractedImage] = []
        sections: List[DocumentSection] = []

        try:
            with zipfile.ZipFile(file_path, "r") as z:
                # 1. Extract embedded images
                media_files = [f for f in z.namelist() if f.startswith("ppt/media/")]
                img_counter = 1
                for mf in media_files:
                    filename = os.path.basename(mf)
                    dest_path = os.path.join(images_dir, filename)
                    with open(dest_path, "wb") as f_out:
                        f_out.write(z.read(mf))

                    rel_path = f"storage/projects/{project_id}/images/{filename}"
                    img_id = f"img-{img_counter:02d}"
                    img_format = os.path.splitext(filename)[1].lower().lstrip(".")
                    extracted_images.append(
                        ExtractedImage(
                            id=img_id,
                            filePath=rel_path,
                            originalName=filename,
                            format=img_format if img_format in ["png", "jpg", "jpeg", "webp", "svg"] else "png"
                        )
                    )
                    img_counter += 1

                # 2. Extract slides
                slide_files = [f for f in z.namelist() if f.startswith("ppt/slides/slide") and f.endswith(".xml")]

                def extract_num(name):
                    m = re.search(r"slide(\d+)\.xml", name)
                    return int(m.group(1)) if m else 9999

                slide_files.sort(key=extract_num)
                title = "Bài giảng PPTX"

                for order, sf in enumerate(slide_files):
                    s_num = extract_num(sf)
                    xml_content = z.read(sf)
                    tree = ET.fromstring(xml_content)

                    lines = []
                    for p_node in tree.iter():
                        if p_node.tag.endswith("}p"):
                            p_text = "".join([t.text for t in p_node.iter() if t.text and t.text.strip()])
                            if p_text.strip():
                                lines.append(re.sub(r"\s+", " ", p_text.strip()))

                    # Extract notes
                    note_file = f"ppt/notesSlides/notesSlide{s_num}.xml"
                    notes = ""
                    if note_file in z.namelist():
                        try:
                            note_tree = ET.fromstring(z.read(note_file))
                            note_lines = []
                            for np_node in note_tree.iter():
                                if np_node.tag.endswith("}p"):
                                    n_text = "".join([t.text for t in np_node.iter() if t.text and t.text.strip()])
                                    if n_text.strip():
                                        note_lines.append(n_text.strip())
                            notes = " ".join(note_lines)
                        except Exception:
                            notes = ""

                    sec_title = lines[0] if lines else f"Slide {s_num}"
                    bullets = lines[1:] if len(lines) > 1 else []
                    if order == 0 and lines:
                        title = lines[0]

                    # Assign image ids
                    sec_img_ids = [img.id for img in extracted_images[:2]] if order == 0 else []

                    sections.append(
                        DocumentSection(
                            sectionId=f"sec-{order+1:02d}",
                            title=sec_title,
                            contentOrder=order,
                            bullets=bullets,
                            tables=[],
                            speakerNotes=notes if notes else None,
                            imageIds=sec_img_ids,
                        )
                    )

            return DocumentContext(
                title=title,
                sourceFileName=os.path.basename(file_path),
                fileType=FileType.PPTX,
                sections=sections,
                extractedImages=extracted_images,
                metadata={"totalSlidesOriginal": len(sections)}
            )
        except zipfile.BadZipFile as e:
            raise ParseFailureError(f"Invalid or corrupted PPTX file: {e}") from e

    # ==========================================
    # 2. Markdown Parser
    # ==========================================
    def _parse_markdown(self, file_path: str, project_id: str, images_dir: str) -> DocumentContext:
        with open(file_path, "r", encoding="utf-8") as f:
            content = f.read()

        lines = content.split("\n")
        title = "Tài liệu bài giảng Markdown"
        author = None
        sections: List[DocumentSection] = []

        current_sec_title = ""
        current_bullets: List[str] = []
        current_tables: List[DocumentTable] = []
        current_notes = []
        in_table = False
        table_headers: List[str] = []
        table_rows: List[List[str]] = []
        content_order = 0

        def flush_section():
            nonlocal current_sec_title, current_bullets, current_tables, current_notes, content_order
            if current_sec_title or current_bullets:
                speaker_notes_str = " ".join(current_notes).strip() if current_notes else None
                sections.append(
                    DocumentSection(
                        sectionId=f"sec-{content_order+1:02d}",
                        title=current_sec_title if current_sec_title else f"Phần {content_order+1}",
                        contentOrder=content_order,
                        bullets=list(current_bullets),
                        tables=list(current_tables),
                        speakerNotes=speaker_notes_str,
                    )
                )
                content_order += 1
                current_bullets = []
                current_tables = []
                current_notes = []

        for line in lines:
            line_str = line.strip()
            if not line_str:
                if in_table:
                    if table_headers:
                        current_tables.append(DocumentTable(headers=table_headers, rows=table_rows))
                    in_table = False
                    table_headers = []
                    table_rows = []
                continue

            # Main Title (# Title)
            if line_str.startswith("# ") and not line_str.startswith("## "):
                title = line_str[2:].strip()
                continue

            # Author detection
            if line_str.lower().startswith("tác giả:") or line_str.lower().startswith("author:"):
                author = line_str.split(":", 1)[1].strip()
                continue

            # Section Header (## Section)
            if line_str.startswith("## ") or line_str.startswith("### "):
                if in_table and table_headers:
                    current_tables.append(DocumentTable(headers=table_headers, rows=table_rows))
                    in_table = False
                    table_headers = []
                    table_rows = []

                flush_section()
                current_sec_title = re.sub(r"^#+\s*", "", line_str)
                continue

            # Markdown Table detection
            if line_str.startswith("|") and line_str.endswith("|"):
                cells = [c.strip() for c in line_str.split("|")[1:-1]]
                if all(re.match(r"^:?-+:?$", c) for c in cells):
                    # Separator line
                    continue
                if not in_table:
                    in_table = True
                    table_headers = cells
                    table_rows = []
                else:
                    table_rows.append(cells)
                continue
            else:
                if in_table:
                    if table_headers:
                        current_tables.append(DocumentTable(headers=table_headers, rows=table_rows))
                    in_table = False
                    table_headers = []
                    table_rows = []

            # Speaker Notes detection
            note_match = re.match(r"^(?:Ghi chú|Note|Speaker note)s?:\s*(.*)", line_str, re.IGNORECASE)
            if note_match:
                current_notes.append(note_match.group(1))
                continue

            # Bullets or Paragraphs
            if re.match(r"^[-*•]\s+", line_str):
                bullet_text = re.sub(r"^[-*•]\s+", "", line_str)
                current_bullets.append(bullet_text)
            else:
                current_bullets.append(line_str)

        # Flush final section
        if in_table and table_headers:
            current_tables.append(DocumentTable(headers=table_headers, rows=table_rows))
        flush_section()

        return DocumentContext(
            title=title,
            author=author,
            sourceFileName=os.path.basename(file_path),
            fileType=FileType.MD,
            sections=sections,
            extractedImages=[],
            metadata={"lineCount": len(lines)}
        )

    # ==========================================
    # 3. DOCX Parser
    # ==========================================
    def _parse_docx(self, file_path: str, project_id: str, images_dir: str) -> DocumentContext:
        try:
            import docx
            from docx.text.paragraph import Paragraph
            from docx.table import Table
            doc = docx.Document(file_path)
        except Exception as e:
            raise ParseFailureError(f"Corrupted or invalid DOCX document: {e}") from e

        title = "Tài liệu bài giảng DOCX"
        sections: List[DocumentSection] = []
        current_sec_title = ""
        current_bullets: List[str] = []
        current_tables: List[DocumentTable] = []
        current_notes = []
        content_order = 0

        def flush_section():
            nonlocal current_sec_title, current_bullets, current_tables, current_notes, content_order
            if current_sec_title or current_bullets or current_tables:
                speaker_notes_str = " ".join(current_notes).strip() if current_notes else None
                sections.append(
                    DocumentSection(
                        sectionId=f"sec-{content_order+1:02d}",
                        title=current_sec_title if current_sec_title else f"Phần {content_order+1}",
                        contentOrder=content_order,
                        bullets=list(current_bullets),
                        tables=list(current_tables),
                        speakerNotes=speaker_notes_str,
                    )
                )
                content_order += 1
                current_bullets = []
                current_tables = []
                current_notes = []

        # Iterate body children in sequential order
        for child in doc.element.body:
            tag = child.tag.split("}")[-1] if "}" in child.tag else child.tag
            if tag == "p":
                p = Paragraph(child, doc)
                text = p.text.strip()
                if not text:
                    continue

                style_name = p.style.name.lower() if p.style else ""
                if "title" in style_name:
                    title = text
                    continue
                elif "heading 1" in style_name or "heading 2" in style_name:
                    flush_section()
                    current_sec_title = text
                    continue

                note_match = re.match(r"^(?:Ghi chú|Note|Speaker note)s?:\s*(.*)", text, re.IGNORECASE)
                if note_match:
                    current_notes.append(note_match.group(1))
                else:
                    current_bullets.append(text)
            elif tag == "tbl":
                tbl = Table(child, doc)
                if tbl.rows:
                    headers = [cell.text.strip() for cell in tbl.rows[0].cells]
                    rows = []
                    for row in tbl.rows[1:]:
                        rows.append([cell.text.strip() for cell in row.cells])
                    current_tables.append(DocumentTable(headers=headers, rows=rows))

        flush_section()

        return DocumentContext(
            title=title,
            sourceFileName=os.path.basename(file_path),
            fileType=FileType.DOCX,
            sections=sections,
            extractedImages=[],
            metadata={"paragraphCount": len(doc.paragraphs), "tableCount": len(doc.tables)}
        )

    # ==========================================
    # ==========================================
    # 4. PDF Parser (Supports both native text & AI Multimodal for scanned PDFs)
    # ==========================================
    def _parse_scanned_pdf_with_gemini(
        self, file_path: str, project_id: str, images_dir: str, doc
    ) -> Optional[DocumentContext]:
        api_key = os.environ.get("GEMINI_API_KEY") or os.environ.get("GOOGLE_API_KEY")
        if not api_key:
            return None

        try:
            from google import genai
            from google.genai import types

            client = genai.Client(api_key=api_key)
            with open(file_path, "rb") as f:
                pdf_bytes = f.read()

            prompt = """Bạn là chuyên gia phân tích tài liệu sư phạm hàng đầu.
Hãy phân tích tài liệu PDF (dạng giáo trình/sách giáo khoa scan hoặc tài liệu bài giảng) đính kèm.
Trích xuất thông tin chi tiết dưới dạng JSON theo đúng schema sau:
{
  "title": "Tiêu đề chính xác của bài giảng (ví dụ: Bài 16: Bạn có thường đến thư viện không? 《你常去图书馆吗》)",
  "subject": "Môn học / Lĩnh vực (ví dụ: Tiếng Trung / Hán ngữ, Ngoại ngữ, Khoa học xã hội...)",
  "author": "Đơn vị/Tác giả (nếu có, không có thì để null)",
  "sections": [
    {
      "title": "Tiêu đề phần học (ví dụ: Từ Mới (生词 - Shēngcí), Ngữ Pháp Trọng Điểm, Bài Khóa...)",
      "bullets": [
        "Nội dung chi tiết từng từ vựng, phiên âm pinyin, từ loại, nghĩa tiếng Việt hoặc quy tắc ngữ pháp kèm ví dụ song ngữ"
      ],
      "speakerNotes": "Lời giảng gợi ý cho giảng viên"
    }
  ]
}
Chỉ trả về chuỗi JSON hợp lệ."""

            config = types.GenerateContentConfig(
                response_mime_type="application/json",
                temperature=0.2,
            )

            models_to_try = ["gemini-flash-lite-latest", "gemini-flash-latest", "gemini-3.8-flash"]
            raw_text = None
            for m in models_to_try:
                try:
                    resp = client.models.generate_content(
                        model=m,
                        contents=[
                            types.Part.from_bytes(data=pdf_bytes, mime_type="application/pdf"),
                            prompt
                        ],
                        config=config
                    )
                    if resp.text:
                        raw_text = resp.text.strip()
                        break
                except Exception as ex:
                    print(f"[WARN] Gemini multimodal model {m} failed: {ex}")

            if not raw_text:
                return None

            if raw_text.startswith("```"):
                lines = raw_text.split("\n")
                if lines[0].startswith("```"):
                    lines = lines[1:]
                if lines and lines[-1].strip() == "```":
                    lines = lines[:-1]
                raw_text = "\n".join(lines).strip()

            parsed = json.loads(raw_text)
            title = parsed.get("title") or "Tài liệu bài giảng PDF"
            subject = parsed.get("subject")
            author = parsed.get("author")

            # Extract page images to provide visual assets for slides
            extracted_images: List[ExtractedImage] = []
            for page_idx in range(min(len(doc), 10)):
                pix = doc[page_idx].get_pixmap(dpi=150)
                img_name = f"page_{page_idx+1:02d}.png"
                img_dest = os.path.join(images_dir, img_name)
                pix.save(img_dest)
                extracted_images.append(
                    ExtractedImage(
                        id=f"img-{page_idx+1:02d}",
                        filePath=f"storage/projects/{project_id}/images/{img_name}",
                        originalName=img_name,
                        format="png"
                    )
                )

            sections: List[DocumentSection] = []
            for s_idx, sec_data in enumerate(parsed.get("sections") or []):
                sections.append(
                    DocumentSection(
                        sectionId=f"sec-{s_idx+1:02d}",
                        title=sec_data.get("title") or f"Phần {s_idx+1}",
                        contentOrder=s_idx,
                        bullets=sec_data.get("bullets") or [],
                        tables=[],
                        speakerNotes=sec_data.get("speakerNotes")
                    )
                )

            return DocumentContext(
                title=title,
                subject=subject,
                author=author,
                sourceFileName=os.path.basename(file_path),
                fileType=FileType.PDF,
                sections=sections,
                extractedImages=extracted_images,
                metadata={"pageCount": len(doc), "isScannedPdf": True, "parsedVia": "GeminiMultimodal"}
            )
        except Exception as e:
            print(f"[WARN] _parse_scanned_pdf_with_gemini encountered error: {e}")
            return None

    def _parse_pdf(self, file_path: str, project_id: str, images_dir: str) -> DocumentContext:
        try:
            import pymupdf
            doc = pymupdf.open(file_path)
        except Exception as e:
            raise ParseFailureError(f"Corrupted or invalid PDF document: {e}") from e

        total_text_len = sum(len(doc[i].get_text("text").strip()) for i in range(len(doc)))

        # If scanned document (no embedded text layer) or very low text, use Gemini Multimodal
        if total_text_len < 60:
            ai_context = self._parse_scanned_pdf_with_gemini(file_path, project_id, images_dir, doc)
            if ai_context and ai_context.sections:
                return ai_context

        # Regular text-based PDF fallback
        title = "Tài liệu bài giảng PDF"
        sections: List[DocumentSection] = []
        extracted_images: List[ExtractedImage] = []

        # Extract page images
        for page_idx in range(min(len(doc), 10)):
            pix = doc[page_idx].get_pixmap(dpi=150)
            img_name = f"page_{page_idx+1:02d}.png"
            img_dest = os.path.join(images_dir, img_name)
            pix.save(img_dest)
            extracted_images.append(
                ExtractedImage(
                    id=f"img-{page_idx+1:02d}",
                    filePath=f"storage/projects/{project_id}/images/{img_name}",
                    originalName=img_name,
                    format="png"
                )
            )

        for page_idx in range(len(doc)):
            page = doc[page_idx]
            text = page.get_text("text").strip()
            if not text:
                continue

            lines = [line.strip() for line in text.split("\n") if line.strip()]
            sec_title = lines[0] if lines else f"Trang {page_idx+1}"
            bullets = []
            notes = []

            for line in lines[1:]:
                note_match = re.match(r"^(?:Ghi chú|Note|Speaker note)s?:\s*(.*)", line, re.IGNORECASE)
                if note_match:
                    notes.append(note_match.group(1))
                else:
                    bullets.append(re.sub(r"^[-*•]\s*", "", line))

            if page_idx == 0 and lines:
                title = lines[0]

            sections.append(
                DocumentSection(
                    sectionId=f"sec-{page_idx+1:02d}",
                    title=sec_title,
                    contentOrder=page_idx,
                    bullets=bullets,
                    tables=[],
                    speakerNotes=" ".join(notes) if notes else None,
                )
            )

        return DocumentContext(
            title=title,
            sourceFileName=os.path.basename(file_path),
            fileType=FileType.PDF,
            sections=sections,
            extractedImages=extracted_images,
            metadata={"pageCount": len(doc)}
        )
