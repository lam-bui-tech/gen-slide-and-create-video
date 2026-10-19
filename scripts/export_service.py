"""
scripts/export_service.py
Multi-format Export Pipeline for AI Lecture Studio:
- MP4: subprocess calling `npx hyperframes render` (1080p, audio, timeout, log capture, cleanup)
- PDF: Headless Chrome print of scene HTML, vector text, Vietnamese fonts embedded
- PPTX: Image-pack (one high-res 1080p PNG per slide + speaker notes, documented as non-editable)
"""

import os
import sys
import json
import sqlite3
import subprocess
import threading
import uuid
from datetime import datetime
from typing import Optional, Dict, Any, Callable
from PIL import Image, ImageDraw, ImageFont
import pptx
from pptx.util import Inches

from schemas.models import LessonSceneJSON, SlideScene

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
STORAGE_DIR = os.path.join(BASE_DIR, "storage", "projects")
DB_PATH = os.path.join(BASE_DIR, "storage", "lecture_studio.db")


def find_chrome_executable() -> str:
    """Finds Google Chrome executable for headless printing across OS platforms."""
    possible_paths = [
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
        r"C:\Users\OS\AppData\Local\Google\Chrome\Application\chrome.exe",
        "google-chrome",
        "chrome",
        "chromium",
        "chromium-browser",
    ]
    for p in possible_paths:
        if os.path.isabs(p) and os.path.exists(p):
            return p
    # Fallback to PATH
    return "chrome"


def measure_video_duration(video_path: str) -> float:
    """Measures video duration in seconds using ffprobe."""
    cmd = [
        "ffprobe",
        "-v",
        "error",
        "-show_entries",
        "format=duration",
        "-of",
        "default=noprint_wrappers=1:nokey=1",
        os.path.abspath(video_path),
    ]
    res = subprocess.run(cmd, capture_output=True, text=True, check=True)
    return float(res.stdout.strip())


def get_vietnamese_font(size: int = 36) -> ImageFont.FreeTypeFont:
    """Returns a Truetype font supporting Vietnamese diacritics."""
    font_candidates = [
        r"C:\Windows\Fonts\segoeui.ttf",
        r"C:\Windows\Fonts\arial.ttf",
        r"C:\Windows\Fonts\tahoma.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
        "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    ]
    for path in font_candidates:
        if os.path.exists(path):
            try:
                return ImageFont.truetype(path, size)
            except Exception:
                pass
    return ImageFont.load_default()


# ==========================================
# 1. PDF Vector Exporter (Headless Chrome)
# ==========================================

def generate_scene_pdf_html(scene: LessonSceneJSON) -> str:
    """Generates print-optimized 16:9 HTML for Chrome headless print with embedded Vietnamese fonts."""
    slides_html = []
    for idx, slide in enumerate(scene.slides):
        elements_html = []
        for el in slide.elements:
            content_str = str(el.content) if el.content is not None else ""
            elements_html.append(f'<div class="slide-element"><p>{content_str}</p></div>')

        slide_block = f"""
        <div class="slide-page">
            <div class="slide-header">
                <span class="badge">Slide {idx + 1:02d} • {slide.templateId}</span>
                <span class="project-title">{scene.title}</span>
            </div>
            <div class="slide-body">
                <h1 class="slide-title">{slide.title}</h1>
                <div class="slide-content">
                    {''.join(elements_html)}
                </div>
            </div>
            <div class="slide-footer">
                <span>AI Lecture Studio • Academic Pro</span>
                <span>{slide.durationSeconds:.1f}s</span>
            </div>
        </div>
        """
        slides_html.append(slide_block)

    full_html = f"""<!DOCTYPE html>
<html lang="vi">
<head>
<meta charset="utf-8">
<title>{scene.title} - PDF Export</title>
<style>
  @import url('https://fonts.googleapis.com/css2?family=Be+Vietnam+Pro:wght@400;500;600;700;800&display=swap');
  @page {{
    size: 1920px 1080px;
    margin: 0;
  }}
  * {{
    box-sizing: border-box;
    margin: 0;
    padding: 0;
  }}
  body {{
    background: #0f172a;
    font-family: 'Be Vietnam Pro', 'Segoe UI', Tahoma, Arial, sans-serif;
    color: #f8fafc;
    -webkit-print-color-adjust: exact;
    print-color-adjust: exact;
  }}
  .slide-page {{
    width: 1920px;
    height: 1080px;
    page-break-after: always;
    break-after: page;
    display: flex;
    flex-direction: column;
    justify-content: space-between;
    padding: 80px 100px;
    position: relative;
    background: linear-gradient(135deg, #090d16 0%, #111827 50%, #0f172a 100%);
    border-bottom: 2px solid #1e293b;
    overflow: hidden;
  }}
  .slide-header {{
    display: flex;
    justify-content: space-between;
    align-items: center;
    border-bottom: 1px solid rgba(255, 255, 255, 0.1);
    padding-bottom: 20px;
  }}
  .badge {{
    background: #0369a1;
    color: #ffffff;
    font-size: 20px;
    font-weight: 600;
    padding: 8px 18px;
    border-radius: 9999px;
    text-transform: uppercase;
    letter-spacing: 0.05em;
  }}
  .project-title {{
    font-size: 24px;
    color: #94a3b8;
    font-weight: 500;
  }}
  .slide-body {{
    flex: 1;
    display: flex;
    flex-direction: column;
    justify-content: center;
    padding: 40px 0;
  }}
  .slide-title {{
    font-size: 56px;
    font-weight: 700;
    color: #38bdf8;
    line-height: 1.25;
    margin-bottom: 36px;
  }}
  .slide-content {{
    font-size: 32px;
    line-height: 1.6;
    color: #e2e8f0;
  }}
  .slide-element {{
    background: rgba(30, 41, 59, 0.7);
    border: 1px solid rgba(56, 189, 248, 0.2);
    border-radius: 16px;
    padding: 30px 40px;
    margin-bottom: 24px;
  }}
  .slide-footer {{
    display: flex;
    justify-content: space-between;
    font-size: 22px;
    color: #64748b;
    border-top: 1px solid rgba(255, 255, 255, 0.1);
    padding-top: 20px;
  }}
</style>
</head>
<body>
{''.join(slides_html)}
</body>
</html>"""
    return full_html


def render_pdf_from_scene(scene: LessonSceneJSON, output_pdf_path: str) -> str:
    """Renders scene HTML to a multi-page vector PDF via Chrome Headless."""
    os.makedirs(os.path.dirname(os.path.abspath(output_pdf_path)), exist_ok=True)
    temp_html_path = output_pdf_path + ".temp.html"

    html_content = generate_scene_pdf_html(scene)
    with open(temp_html_path, "w", encoding="utf-8") as f:
        f.write(html_content)

    chrome_bin = find_chrome_executable()
    abs_html = os.path.abspath(temp_html_path).replace("\\", "/")
    abs_pdf = os.path.abspath(output_pdf_path)

    cmd = [
        chrome_bin,
        "--headless=new",
        "--disable-gpu",
        "--no-sandbox",
        "--allow-file-access-from-files",
        f"--print-to-pdf={abs_pdf}",
        "--no-pdf-header-footer",
        abs_html,
    ]

    try:
        res = subprocess.run(cmd, capture_output=True, text=True, timeout=60)
        if res.returncode != 0 or not os.path.exists(abs_pdf):
            raise RuntimeError(f"Chrome PDF render failed: code {res.returncode}, stderr: {res.stderr}")
        return abs_pdf
    finally:
        if os.path.exists(temp_html_path):
            try:
                os.remove(temp_html_path)
            except Exception:
                pass


# ==========================================
# 2. PPTX Image-Pack Exporter
# ==========================================

def render_pptx_image_pack(scene: LessonSceneJSON, output_pptx_path: str) -> str:
    """
    Renders scene to PPTX image-pack presentation:
    - 16:9 widescreen layout
    - 1 high-resolution 1080p PNG per slide
    - Speaker notes containing Vietnamese narration and user notes
    - Documented clearly as non-editable
    """
    os.makedirs(os.path.dirname(os.path.abspath(output_pptx_path)), exist_ok=True)
    temp_dir = output_pptx_path + "_frames"
    os.makedirs(temp_dir, exist_ok=True)

    prs = pptx.Presentation()
    # 16:9 Widescreen standard in inches
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    title_font = get_vietnamese_font(48)
    body_font = get_vietnamese_font(30)
    badge_font = get_vietnamese_font(22)

    try:
        for idx, slide in enumerate(scene.slides):
            # 1. Render high-res PNG frame (1920x1080)
            img = Image.new("RGB", (1920, 1080), color="#0f172a")
            draw = ImageDraw.Draw(img)

            # Background subtle card
            draw.rectangle([80, 80, 1840, 1000], fill="#1e293b", outline="#0284c7", width=3)

            # Header info
            badge_text = f"SLIDE {idx + 1:02d} • {slide.templateId.upper()}"
            draw.text((120, 120), badge_text, fill="#38bdf8", font=badge_font)

            # Slide Title
            draw.text((120, 180), slide.title, fill="#ffffff", font=title_font)

            # Content elements
            y_offset = 320
            for el in slide.elements:
                content_str = str(el.content) if el.content is not None else ""
                # Draw lines
                draw.text((120, y_offset), f"• {content_str}", fill="#cbd5e1", font=body_font)
                y_offset += 80

            # Footer
            draw.text((120, 930), f"AI Lecture Studio • {scene.title}", fill="#64748b", font=badge_font)

            frame_img_path = os.path.join(temp_dir, f"slide_{idx:03d}.png")
            img.save(frame_img_path, format="PNG")

            # 2. Add slide to presentation and insert full-screen image
            ppt_slide = prs.slides.add_slide(blank_layout)
            ppt_slide.shapes.add_picture(
                frame_img_path,
                0,
                0,
                width=prs.slide_width,
                height=prs.slide_height,
            )

            # 3. Add Speaker Notes with Vietnamese Narration and Disclaimer
            notes_slide = ppt_slide.notes_slide
            tf = notes_slide.notes_text_frame

            disclaimer = (
                "[IMAGE-PACK EXPORT]: This presentation contains high-resolution slide renders with speaker notes. "
                "Visual layout is preserved exactly as rendered in AI Lecture Studio and is not directly editable."
            )
            narration_text = slide.voice.narration if slide.voice and slide.voice.narration else "(No narration)"
            notes_content = slide.notes if slide.notes else ""

            tf.text = (
                f"{disclaimer}\n\n"
                f"--- NARRATION (THUYẾT MINH) ---\n"
                f"{narration_text}\n\n"
                f"--- SPEAKER NOTES (GHI CHÚ GIẢNG VIÊN) ---\n"
                f"{notes_content}"
            )

        abs_pptx = os.path.abspath(output_pptx_path)
        prs.save(abs_pptx)
        return abs_pptx
    finally:
        # Cleanup temp frames
        if os.path.exists(temp_dir):
            import shutil
            try:
                shutil.rmtree(temp_dir)
            except Exception:
                pass


# ==========================================
# 3. MP4 HyperFrames Exporter (Subprocess)
# ==========================================

def render_mp4_hyperframes(
    project_dir: str,
    output_mp4_path: str,
    timeout_seconds: int = 300,
    log_callback: Optional[Callable[[str], None]] = None,
) -> str:
    """
    Renders MP4 video using `npx hyperframes render`:
    - 1080p resolution with audio
    - Subprocess with timeout and log capture
    - Error handling and status tracking
    """
    os.makedirs(os.path.dirname(os.path.abspath(output_mp4_path)), exist_ok=True)
    abs_out = os.path.abspath(output_mp4_path)
    abs_proj = os.path.abspath(project_dir)

    # Flags verified from `npx hyperframes render --help`:
    # npx hyperframes render [DIR] -o <output> --resolution=1080p -f 30 -q looks
    cmd = [
        "npx",
        "hyperframes",
        "render",
        abs_proj,
        "-o",
        abs_out,
        "--resolution=1080p",
        "-f",
        "30",
        "-q",
        "looks",
    ]

    # Run via shell on Windows
    use_shell = sys.platform == "win32"
    process = subprocess.Popen(
        cmd,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        shell=use_shell,
        cwd=BASE_DIR,
    )

    logs = []
    try:
        stdout, _ = process.communicate(timeout=timeout_seconds)
        if stdout:
            logs.append(stdout)
            if log_callback:
                log_callback(stdout)

        if process.returncode != 0:
            raise RuntimeError(f"HyperFrames render failed with code {process.returncode}:\n{stdout}")

        if not os.path.exists(abs_out) or os.path.getsize(abs_out) == 0:
            raise RuntimeError("HyperFrames render completed but output MP4 is missing or empty.")

        return abs_out
    except subprocess.TimeoutExpired:
        process.kill()
        raise TimeoutError(f"HyperFrames render timed out after {timeout_seconds} seconds.")


# ==========================================
# 4. Export Service & Job Coordinator
# ==========================================

class ExportService:
    """Manages asynchronous export jobs for MP4, PDF, and PPTX."""

    def __init__(self, db_path: str = DB_PATH, storage_base: str = STORAGE_DIR):
        self.db_path = db_path
        self.storage_base = storage_base

    def create_job(self, project_id: str, export_type: str) -> Dict[str, Any]:
        """Creates a pending export job in SQLite."""
        job_id = f"exp-{uuid.uuid4().hex[:8]}"
        created_at = datetime.utcnow().isoformat() + "Z"

        with sqlite3.connect(self.db_path) as conn:
            cursor = conn.cursor()
            cursor.execute(
                """
                INSERT INTO jobs (id, project_id, type, status, progress, download_url, error)
                VALUES (?, ?, ?, 'pending', 0, NULL, NULL)
                """,
                (job_id, project_id, export_type),
            )
            conn.commit()

        return {
            "exportId": job_id,
            "projectId": project_id,
            "type": export_type,
            "status": "pending",
            "createdAt": created_at,
        }

    def get_job(self, job_id: str) -> Optional[Dict[str, Any]]:
        """Retrieves job status from SQLite."""
        with sqlite3.connect(self.db_path) as conn:
            cursor = conn.cursor()
            cursor.execute(
                """
                SELECT id, project_id, type, status, progress, download_url, error
                FROM jobs WHERE id = ?
                """,
                (job_id,),
            )
            row = cursor.fetchone()

        if not row:
            return None

        return {
            "exportId": row[0],
            "projectId": row[1],
            "type": row[2],
            "status": row[3],
            "progress": row[4],
            "downloadUrl": row[5],
            "error": row[6],
        }

    def update_job(
        self,
        job_id: str,
        status: str,
        progress: int = 0,
        download_url: Optional[str] = None,
        error: Optional[str] = None,
    ):
        """Updates job status, progress, download_url, or error."""
        with sqlite3.connect(self.db_path) as conn:
            cursor = conn.cursor()
            cursor.execute(
                """
                UPDATE jobs
                SET status = ?, progress = ?, download_url = ?, error = ?
                WHERE id = ?
                """,
                (status, progress, download_url, error, job_id),
            )
            conn.commit()

    def render_mp4_with_cleanup(
        self,
        project_id: str,
        scene: LessonSceneJSON,
        output_path: str,
        render_executor: Optional[Callable] = None,
    ) -> str:
        """Renders MP4 with cleanup guarantee if failure happens."""
        executor = render_executor or render_mp4_hyperframes
        try:
            proj_dir = os.path.join(self.storage_base, project_id)
            return executor(proj_dir, output_path)
        except Exception as e:
            if os.path.exists(output_path):
                try:
                    os.remove(output_path)
                except Exception:
                    pass
            raise e

    def execute_job_sync(self, job_id: str):
        """Executes the export job synchronously (called in background worker)."""
        job = self.get_job(job_id)
        if not job:
            return

        project_id = job["projectId"]
        export_type = job["type"].lower()

        # Update to running
        self.update_job(job_id, status="running", progress=10)

        try:
            # Load project scene
            scene_path = os.path.join(self.storage_base, project_id, "scene.json")
            if not os.path.exists(scene_path):
                raise FileNotFoundError(f"scene.json not found for project {project_id}")

            with open(scene_path, "r", encoding="utf-8") as f:
                scene_data = json.load(f)
            scene = LessonSceneJSON.model_validate(scene_data)

            exports_dir = os.path.join(self.storage_base, project_id, "exports")
            os.makedirs(exports_dir, exist_ok=True)

            self.update_job(job_id, status="running", progress=35)

            if export_type == "pdf":
                out_path = os.path.join(exports_dir, "lecture.pdf")
                rel_url = f"/storage/projects/{project_id}/exports/lecture.pdf"
                render_pdf_from_scene(scene, out_path)
            elif export_type == "pptx":
                out_path = os.path.join(exports_dir, "lecture.pptx")
                rel_url = f"/storage/projects/{project_id}/exports/lecture.pptx"
                render_pptx_image_pack(scene, out_path)
            elif export_type == "mp4":
                out_path = os.path.join(exports_dir, "lecture_1080p.mp4")
                rel_url = f"/storage/projects/{project_id}/exports/lecture_1080p.mp4"
                self.render_mp4_with_cleanup(project_id, scene, out_path)
            else:
                raise ValueError(f"Unsupported export type: {export_type}")

            self.update_job(job_id, status="completed", progress=100, download_url=rel_url)
        except Exception as e:
            self.update_job(job_id, status="failed", progress=0, error=str(e))

    def trigger_background_export(self, job_id: str):
        """Dispatches job execution to a background worker thread."""
        worker = threading.Thread(target=self.execute_job_sync, args=(job_id,), daemon=True)
        worker.start()
